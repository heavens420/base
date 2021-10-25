import openpyxl as xl

path_ta = r'C:\Users\420\Desktop\kdgc\6月集团考核\华为-IP-实指令管理模板_0619.xlsx'
path_sa = r'C:\Users\420\Desktop\kdgc\6月集团考核\华为-IP-实指令管理模板_0623.xlsx'
wb1 = xl.load_workbook(path_ta)
wb2 = xl.load_workbook(path_sa)
sheet_names1 = wb1.sheetnames
sheet_names2 = wb2.sheetnames
print(sheet_names1)
ws1 = wb1[sheet_names1[1]]
ws2 = wb2[sheet_names2[0]]

ta_code = ws1['F']
sa_code = ws2['A']

lst = list()
num = 0
for k in range(0, 228):
    value = ta_code[k].value
    if value is not None:
        lst.append(value)
        num += 1
        # print(value)
print(f'num:{num},len:{len(lst)}')

yi = 0
wei = 0
#
# print(len(sa_code))
# print(len(ta_code))
for i in range(len(sa_code)):
    # ta = ta_code[i].value
    sa = sa_code[i].value
    if sa is not None:
        # print(sa)
        if sa in lst:
            print('已核查')
            yi += 1
        else:
            print('未核查')
            wei += 1
print(f'已核查{yi}，未核查{wei}')
